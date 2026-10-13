"""Generate editable course submission reports. Run with the bundled document runtime."""
from pathlib import Path
from copy import deepcopy
import zipfile, hashlib, json, sys
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'01-提交文档';OUT.mkdir(exist_ok=True)
if any(OUT.glob('*.docx')) and '--overwrite' not in sys.argv:
    raise SystemExit('现有文档包含成员修订。请直接编辑 Word；仅在明确需要重建基础草稿时使用 --overwrite。')
MEMBER_NAMES=['陈辰','董宇曦','刘根森','李怡宁','赵毓豪']
N={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
REF=ROOT/'04-原始要求/2026-SECD-文档提交封面模板.docx'

# Surgical cover patch. Every package part except document.xml stays byte-identical.
with zipfile.ZipFile(REF) as z:
    parts={n:z.read(n) for n in z.namelist()}
root=E.fromstring(parts['word/document.xml']);body=root.find('w:body',N)
paras=body.findall('w:p',N);table=body.find('w:tbl',N)
def replace_text(el,value,size=None):
    text_nodes=el.findall('.//w:t',N)
    if text_nodes:
        text_nodes[0].text=value
        for t in text_nodes[1:]:t.text=''
    else:
        p=el if el.tag==qn('w:p') else el.find('w:p',N)
        r=E.SubElement(p,qn('w:r'));E.SubElement(r,qn('w:t')).text=value
    if size:
        for r in el.findall('.//w:r',N):
            rp=r.find('w:rPr',N)
            if rp is None:rp=E.SubElement(r,qn('w:rPr'))
            sz=rp.find('w:sz',N)
            if sz is None:sz=E.SubElement(rp,qn('w:sz'))
            sz.set(qn('w:val'),str(size*2))
replace_text(paras[6],'组号：待填写')
replace_text(paras[7],'小组成员表（组长：陈辰）：')
rows=table.findall('w:tr',N)
for i,row in enumerate(rows[1:6]):
    cells=row.findall('w:tc',N)
    for cell,value in zip(cells,[str(i+1),['25230426','87220604','14230107','35231905','27200221'][i],MEMBER_NAMES[i],'待确认','']):replace_text(cell,value,12)
table.remove(rows[6])
for p in paras[10:]:body.remove(p)
# Remove the empty spacer after the member table to keep the required note on the cover.
body.remove(paras[8])
# Blank source spacer before group number can otherwise push the signature note to a second page.
body.remove(paras[5])
parts['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(OUT/'00-提交封面.docx','w',zipfile.ZIP_DEFLATED) as z:
    for n,data in parts.items():z.writestr(n,data)

TEAMS=[
 ['陈辰','组长与集成','范围管理、架构、合并验收、全组总结','20%（建议）'],
 ['董宇曦','需求与学生端','用例、目录、选课、课表、成绩查询','20%（建议）'],
 ['刘根森','后端与数据库','数据模型、事务、锁、容量和先修校验','20%（建议）'],
 ['李怡宁','教师与教务端','授课、名单、成绩录入、人员维护','20%（建议）'],
 ['赵毓豪','测试与交付','异常与并发测试、部署、证据、操作说明','20%（建议）'],
]
DOCS=[]
def doc(name,title,pages):DOCS.append((name,title,pages))
def h(t):return ('h',t)
def p(t):return ('p',t)
def t(headers,rows):return ('t',headers,rows)
def bullets(*items):return ('b',items)

doc('01-开发计划','课程注册系统开发计划',[
 [p('本计划覆盖高校课程注册系统的实现、测试和课程材料交付。采用 Vue 3、Flask、PostgreSQL 和 Docker Compose，优先完成题目中的八类核心用例。项目由五位成员协作，陈辰 为组长。'),
 h('1 项目范围与交付物'),p('必交内容包括可运行程序、开发计划、需求分析、设计文档、测试计划、测试报告、五份个人工作总结和组长汇总报告。本项目同时提供配置管理计划、初始化数据、数据库结构 SQL、源代码和运行说明。'),
 p('旧课程目录及计费系统通过独立 HTTP 模拟服务对接。真实 Ingres、DEC VAX 和生产财务系统不在本次实现环境内。课程 PPT 允许 Web 技术，故采用浏览器界面替代题目中的 Windows 95/98 桌面界面。'),
 h('2 人员安排'),t(['成员','职责','对应工作','拟定比例'],TEAMS),p('以上为建议分工和工作量安排，不等同于真实贡献记录。各成员需按实际承担的工作修订个人总结，组长核实后将贡献比例填写到封面，合计必须为 100%。'),
 h('3 协作方式'),p('以功能模块交付为单位，需求编号贯穿设计和测试。成员提交修改前运行相关测试；陈辰 负责审查接口、组织集成和维护版本清单。每日记录完成内容、问题及下一步计划，发现范围变化时先修订需求与测试用例。')],
 [h('4 进度与迭代计划'),p('题目给出的课程周期为 2026 年 9 月 8 日至 10 月 16 日。本计划以 10 月 13 日当前版本为起点安排剩余工作，不追溯编造此前开发记录。'),
 t(['迭代','计划时间','目标','验收条件'],[
 ['I1','10 月 13 日','核心功能、数据与容器启动','三类角色可登录，核心事务测试通过'],
 ['I2','10 月 14 日','组员走查与问题修复','逐条核对用例及异常流，修订文档'],
 ['I3','10 月 15 日','答辩演练与交付冻结','从干净环境启动，完整演示，无阻断缺陷'],
 ['验收','10 月 16 日','课程运行检查与提交','封面资料与签名齐全，文档源码版本一致']]),
 h('5 风险及应对'),t(['风险','影响','应对'],[
 ['并发选课超额','业务规则失效','对教学班加行锁，真实 PostgreSQL 并发测试'],
 ['外部服务不可用','目录或计费中断','目录缓存和超时；计费持久化重试与幂等'],
 ['环境不一致','无法现场演示','Compose、固定依赖、预构建前端、初始化数据'],
 ['需求歧义','实现与题目不一致','在需求分析中记录解释与验收边界'],
 ['人员信息缺失','不符合提交形式','组长补充组号、学号、真实贡献及本人签名']]),
 h('6 完成标准'),bullets('八类核心用例及关键异常流有对应实现和测试。','系统可通过 Docker Compose 启动，重建容器后数据仍在。','源码、数据库模型、测试证据和文档相互一致。','实际性能结果与尚未验证的性能目标分别列明。','所有成员核对个人总结，组长汇总并完成提交封面。')]
])

doc('02-需求分析文档','课程注册系统需求分析',[
 [h('1 背景与系统边界'),p('Wylie College 需要替换旧有学生注册系统。新系统读取外部课程目录，管理本地师生、授课和选课数据，并向计费系统发送最终课表与学费信息。课程目录由其他系统维护，本系统不提供写入旧课程目录的接口。'),
 h('2 术语与参与者'),t(['术语','定义'],[['Course 课程','课程定义，如数据结构，含课程号、院系、学分、费用及先修要求'],['Course Offering 教学班','某学期某课程的具体开班，含时间、教室、教师和选课人数'],['Schedule 课表','学生一个学期的主选和按顺序排列的备选记录'],['Registrar 教务管理员','管理师生信息并关闭选课的授权角色'],['Report Card 成绩单','学生上一已结束学期所修课程的成绩']]),
 t(['参与者','允许操作'],[['学生','登录、查目录、保存草稿、提交/退改/删除本人课表、查看本人成绩'],['教师','登录、选择具备资格的授课班、查本人授课名单、录入本人班级成绩'],['教务管理员','登录、增改删师生、管理授课资格、关闭选课、查计费和审计记录'],['外部系统','课程目录只读查询；计费服务接收通知并按唯一键去重']]),
 h('3 功能需求编号'),t(['编号','功能','验收要点'],[['FR01','登录与权限','有效账号进入对应工作台；无效账号和越权请求被拒绝'],['FR02','维护师生信息','教务独占维护权，自动生成内部编号，处理不存在和删除取消'],['FR03','课程注册','4 主选 + 2 备选、草稿、提交、退改、删除、异常处理'],['FR04','选择授课','资格、归属和时间冲突校验'],['FR05','提交成绩','仅任课教师对已结束学期名单录入 A/B/C/D/F/I'],['FR06','查看成绩单','仅访问本人上一学期成绩；无记录时明确提示'],['FR07','关闭选课','互斥检查、取消教学班、备选补位、生成最终计费'],['FR08','外部系统交互','只读目录、超时提示、计费重试及幂等']])],
 [h('4 核心用例说明'),
 p('UC01 登录：参与者输入账号和密码。系统校验有效状态和密码摘要，创建会话并返回角色信息。错误账号、错误密码或停用状态统一提示；退出登录清除会话。'),
 p('UC02 维护学生 / UC03 维护教师：教务管理员新增人员，系统生成唯一内部 ID。更新前查询人员，修改姓名、出生日期、身份识别号、状态、院系和毕业日期等适用字段。删除前显示确认框，取消时不执行修改。不存在返回 404；仍关联当前学期业务时拒绝删除，须先解除关联。历史记录通过逻辑删除保留。'),
 p('UC04 注册课程：学生获取本学期教学班，选择 4 门主选及 2 门备选。草稿可不完整且不占名额。首次提交必须满足 4+2；已提交课表允许在开放期保留 1 至 4 门主选、0 至 2 门备选，以支持退改选；退空时使用删除课表。备选顺序即补位顺序，备选不预占名额。'),
 p('提交时核验教学班存在且属于本学期、课程不重复、先修课通过、时间不冲突及名额可用。任一错误导致整个修改回滚，保留旧课表和名额。提交后再修改必须重新提交，不允许将已占名额的课表直接降为草稿。目录不可用时提示重试；选课关闭时拒绝变更。'),
 p('UC05 选择授课：教师查看有授课资格的教学班及本人安排。选择或取消授课后保存，系统拒绝其他教师已占班级、无授课资格和时间冲突。保存失败不改变原安排。关闭选课后不能修改授课。'),
 p('UC06 提交成绩：教师选择已结束学期的本人班级，系统展示实际选课名单及已有成绩。教师可录入或改写 A、B、C、D、F、I，留空表示跳过并保留原成绩。无效等级或非该班学生导致整个批次回滚。'),
 p('UC07 查看成绩：学生查看上一学期课程与成绩；尚未录入显示“待录入”，没有修课记录显示空状态。请求不接受用其他学生编号改变数据范围。'),
 p('UC08 关闭选课：教务确认后尝试取得学期独占锁，若存在进行中的选课或授课修改事务则拒绝并提示稍后重试。取消无教师教学班，对已提交课表按备选顺序补位，最终取消人数不足 3 的教学班。一次事务内保存最终课表、关闭状态和计费待发送记录。重复关闭被拒绝。')],
 [h('5 业务规则和歧义处理'),t(['编号','规则 / 本次解释'],[
 ['BR01','课程与教学班分开建模，同一课程可有多个时间安排，但一份课表不得重复选同一课程。'],
 ['BR02','每班最多 10 人，最终至少 3 人且有教师才开班；只统计 enrolled 名额。'],
 ['BR03','同一天的节次区间相交即冲突；先修课程成绩 A/B/C/D 视为通过。'],
 ['BR04','关闭时以原有达标主选为已确认课程，按课表 ID 顺序处理补位；遇满员、冲突或未通过先修课则尝试下一备选。'],
 ['BR05','未达标主选可以在有有效备选时被替换。所有补位结束后再判断最低人数，允许备选加入使教学班达到最低开班数。'],
 ['BR06','题目概述提及提交后计费，详细 Close Registration 用例要求关闭后按最终课表计费，本实现采用详细用例，避免课程取消后错误计费。'],
 ['BR07','没有有效备选时不强行凑满 4 门。取消课程不计费；最终无课程的已提交课表生成零金额通知。'],
 ['BR08','草稿不参加补位和计费。成绩不删除历史记录。毕业日期仅对学生有业务意义。']]),
 h('6 非功能要求与验收边界'),t(['原题目标','本次实现与验证边界'],[['2000 中央库并发 / 500 本地服务并发','保留为性能目标；当前验证两请求争抢最后一个名额，未进行该规模压力测试'],['旧目录访问不超过 10 秒','设置 3 秒外部请求超时与每进程 5 秒只读缓存；真实旧系统性能未知'],['80% 事务在 2 分钟内完成','保留为目标；不以功能测试代替正式吞吐和延迟测量'],['7×24 小时 / 停机不超过 10%','单机 Docker 演示不构成可用性承诺；未做长期运行统计'],['成绩安全','会话认证、CSRF、后端归属校验、密码摘要、日志不记录成绩内容'],['Windows 桌面要求','依课程 PPT 允许的 Vue + Flask Web 方案实现']]),
 h('7 验收资料'),p('按开发计划进行课程演示，使用测试计划检查正常流程和异常流。测试报告列出本次实测结果，所有性能目标保留“未验证”状态，直至有对应原始测试数据。')]
])

doc('03-设计文档','课程注册系统设计文档',[
 [h('1 总体架构'),p('采用单体 Flask API 与 Vue 单页应用。开发时通过 Vite 代理 /api，生产演示由 Gunicorn 运行 Flask，同时提供 Vue 构建后的静态文件。PostgreSQL 独立容器存储业务数据，worker 处理计费重试，mock 容器模拟旧课程目录和计费服务。所有容器由 Compose 编排。'),
 t(['层 / 服务','职责','关键文件'],[['浏览器 Vue','角色导航、表单、课程筛选、课表、反馈','frontend/src/App.vue、style.css'],['Flask API','认证、校验、用例事务和数据查询','backend/app.py'],['数据模型','关系、唯一性和字段约束','backend/models.py、schema.sql'],['PostgreSQL','持久化业务数据及并发控制','compose.yaml 中 db'],['计费 worker','读取待发送记录，重试、更新状态','backend/worker.py'],['外部模拟服务','只读目录、故障开关、幂等收据','backend/mock_services.py'],['初始化','创建表和演示数据，已有数据库不重置','backend/seed.py、catalog_data.py']]),
 h('2 部署与接口边界'),p('浏览器访问 http://127.0.0.1:8080。数据库开发端口仅绑定 127.0.0.1:55432；API 与模拟服务通过 Docker 内网通信。PostgreSQL 使用命名卷 postgres-data，模拟计费收据使用 mock-data。数据库健康后才初始化，初始化成功后再启动 API 和 worker。'),
 p('系统对旧目录仅发送 GET，不提供课程编辑页面。示例目录为固定数据，初始化时保存同一课程定义的本地快照以供事务校验。实时展示使用只读接口数据，5 秒缓存减少重复请求。真实系统接入时需补充快照刷新策略及版本一致性机制，当前未连接真实 Ingres。'),
 h('3 状态设计'),t(['对象','状态 / 转移'],[['学期','open → closed；关闭后不提供重新开放接口'],['课表','草稿 → 已提交；已提交后退改保持已提交；删除释放名额'],['选项','selected（草稿/备选）→ enrolled（有效选课）；取消班级 → cancelled'],['教学班','open → closed 或 cancelled'],['计费通知','pending → sent；失败仍为 pending，增加次数和下次时间']])],
 [h('4 数据模型与约束'),t(['实体','主键及主要字段','关联 / 约束'],[
 ['User','id、username、password_hash、role、status、name','username 唯一；角色与状态检查约束'],
 ['Term','id、name、state、current','状态 open / closed；初始化配置唯一当前学期'],
 ['Course','id、code、credits、fee_cents、prerequisite_id','code 唯一；先修课自关联'],
 ['Offering','id、course_id、term_id、professor_id、day、start、end','课程/学期/教师外键；节次合法范围'],
 ['Eligibility','professor_id、course_id','复合主键表示教师授课资格'],
 ['Schedule','id、student_id、term_id、submitted','学生 + 学期唯一'],
 ['Choice','id、schedule_id、offering_id、kind、rank、status','课表 + 教学班唯一；删除课表级联选项'],
 ['Grade','student_id、offering_id、value','复合主键；等级检查约束'],
 ['Invoice','id、student_id、term_id、amount_cents、payload、status、attempts','学生 + 学期唯一；保存最终课表快照'],
 ['Audit','id、actor_id、action、detail、created_at','关联操作者；记录变更类型和对象编号']]),
 h('5 实体关系'),bullets('一个学期包含多个教学班；一门课程可在不同学期和时段开设多个教学班。','一个学生在一个学期拥有最多一份课表；课表通过 Choice 关联教学班，形成多对多选课关系。','一个教师可授课多个班级，每个教学班最多关联一个教师；Eligibility 独立表示资格。','学生与教学班共同标识成绩；学生与学期共同标识唯一计费通知。'),
 p('费用以整数“分”存储，避免浮点运算误差。关闭时将费用和最终课程清单存入 Invoice.payload，后续重试使用同一快照，不重新计算。人员采用逻辑删除，以保留成绩与审计外键。'),
 h('6 数据初始化'),p('提供 1 个教务账号、3 个教师账号、14 个学生账号、10 门课程、11 个本学期教学班和 2 个历史教学班。student02 至 student04 已选 4 门核心课程，便于演示开班阈值。student13、student14 无先修成绩，可演示先修校验失败。所有初始数据均为课程演示数据。')],
 [h('7 事务与并发设计'),p('正常选课取得学期共享 advisory lock，再锁定当前学生行，按 ID 升序锁定新旧涉及的教学班，统一重算容量。验证通过后整体替换选项并提交；失败时回滚，避免丢失原有课表。不同学生竞争同一教学班的最后名额时，由 PostgreSQL 行锁串行完成最终校验。'),
 p('教师修改同样取得学期共享锁、本人行锁与教学班行锁，再检查教师归属和冲突。锁顺序固定，减少死锁风险。管理员关闭选课尝试独占学期 advisory lock；若存在共享锁，立即返回 409。若取得锁，后续选课请求会等候关闭事务完成，然后读取 closed 状态并被拒绝。'),
 h('8 关闭选课算法'),bullets('取消没有教师的教学班，并将对应选择标记 cancelled。','对最初有教师且人数不少于 3 的主选保留确认资格。','按课表 ID 升序处理已提交课表，备选按 rank 顺序尝试。对未确认主选进行替换，或填补已取消主选留下的空位。','每次补位检查最多 4 门、教学班最多 10 人、先修与冲突。暂未凑满人数的班级在所有补位后统一判断。','最终不足 3 人的班级及其选项取消。按剩余 enrolled 课程生成唯一 Invoice，学期改为 closed，统一提交。'),
 h('9 外部服务可靠性'),p('计费采用事务性待发送表。worker 每 3 秒扫描到期 pending 记录，使用 FOR UPDATE SKIP LOCKED 分批认领最多 10 条，HTTP 超时为 3 秒。失败后按 6、12、24、48、60 秒等间隔重试，延迟封顶 60 秒，但不限制总次数。'),
 p('幂等键为 course-{term_id}-{student_id}。模拟计费服务将键保存到具有唯一主键的 SQLite 收据表，重复投递不生成第二份收据。若远端成功但本地提交失败，下次重试仍使用同一键。SQLite 仅用于独立模拟服务，实际业务数据库始终是 PostgreSQL。'),
 h('10 安全与异常'),p('密码保存 Werkzeug scrypt 摘要；Cookie 启用 HttpOnly 和 SameSite=Lax，会话有效期两小时。写请求必须携带 X-CSRF-Token，登录后轮换令牌。所有权限在后端判断，学生查询使用会话身份，教师查询验证任课归属。400 表示输入错误，401 未登录，403 权限或令牌错误，404 不存在，409 业务冲突，503 外部服务不可用。')],
 [h('11 HTTP 接口清单'),t(['方法与路径','角色','含义'],[
 ['GET /api/session','所有人','取得登录状态及 CSRF 令牌'],['POST /api/login /logout','所有人 / 已登录','登录、退出'],['GET /api/overview /catalog','已登录','学期概览、课程目录'],['GET /api/schedule','学生','查询本人课表'],['PUT /api/schedule','学生','primary、alternate、submit 保存或提交'],['DELETE /api/schedule','学生','删除本人课表'],['GET /api/report-card','学生','上一学期成绩'],['GET /PUT /api/teaching','教师','可授课和本人教学班 / 更新安排'],['GET /api/offerings/{id}/roster','任课教师','本人教学班名单'],['PUT /api/offerings/{id}/grades','任课教师','批量成绩 [{student_id,value}]'],['GET /POST /api/users','教务','查询、创建师生'],['PUT /DELETE /api/users/{id}','教务','修改 / 逻辑删除人员'],['GET /PUT /api/users/{id}/eligibility','教务','查询 / 设置教师课程资格'],['POST /api/registration/close','教务','confirm:true 关闭选课'],['GET /api/invoices /audit','教务','计费通知、操作日志'],['GET /PUT /api/integrations','教务','外部模拟服务状态']]),
 h('12 页面与接口对应'),p('课程目录提供关键词、上课日和余量筛选，选课清单独立保存主选与备选顺序；周课表只显示 enrolled 记录。教师端将本学期名单查询和历史学期成绩录入放在同一入口。教务端分别展示人员、关闭流程、计费记录、操作日志和外部服务状态。'),
 h('13 可维护性与限制'),p('当前代码针对单个课程设计规模，核心业务集中在 app.py，模型及后台任务独立文件。后续扩大规模时可拆分蓝图和服务层、引入数据库迁移、登录限流、课程目录版本同步及正式监控。当前默认凭据和 HTTP 只用于本地演示，真实部署需更换密钥并启用 HTTPS。')]
])

doc('04-测试计划','课程注册系统测试计划',[
 [h('1 测试目标与环境'),p('验证八类核心用例、权限隔离、异常回滚、课程人数边界和外部接口恢复。自动化测试在独立 PostgreSQL 数据库 course_test 中运行；每个用例重建该测试库中的表并初始化固定数据。测试代码拒绝对其他数据库执行重置，演示库 course 不参与破坏性测试。'),
 p('环境为 Docker Desktop、PostgreSQL 18 Alpine、Python 3.13、Flask 3.1.2、SQLAlchemy 2.0.43、pytest 8.4.2。浏览器检查使用本机内置浏览器。前端用 npm 锁文件安装并执行 Vite 生产构建。'),
 h('2 测试层次与进入条件'),bullets('功能集成测试：通过 Flask 测试客户端请求真实 PostgreSQL，验证响应与数据库最终状态。','并发测试：两个独立请求同时争抢第 10 个名额，预期一成一败。','外部系统联调：独立 HTTP mock 容器，验证故障、恢复及重复请求去重。','界面测试：以三类角色操作主要页面，检查成功提示、列表数据及响应布局。','进入条件：Docker 数据库健康，初始化完成，前端生产构建成功。'),
 h('3 需求覆盖矩阵'),t(['需求','主要检查','自动化定位'],[['FR01','未登录、CSRF、错误密码、角色越权、停用账户','unauthenticated、bad_login、role_boundaries、inactive_login'],['FR02','增改删、唯一 ID、关联删除保护、授课资格','user_lifecycle、delete_current、missing_eligibility'],['FR03','草稿、4+2、去重、先修、冲突、退改、容量','draft、first_submit、duplicate、prerequisite、conflict、last_seat'],['FR04','教师归属、冲突、关闭后禁止修改','teaching_conflict、teaching_update'],['FR05/06','名单归属、合法等级、批量回滚、成绩查询','teacher_roster、grade_edit、invalid_grade、grade_batch'],['FR07','在途事务、取消、补位、计费、重复关闭','close_busy、close_cancels、teacherless、alternate_fallback'],['FR08','目录故障、计费重试、幂等','external_catalog、billing_retry、check_integrations.py']])],
 [h('4 关键测试步骤及预期结果'),t(['场景','输入 / 操作','预期结果'],[['正常选课','student01 主选 2/3/4/5，备选 6/7','4 门主选占位；2 门备选不占位'],['人数边界','已有 9 人，两学生同时提交','一个 200，一个 409，最终 10 人'],['先修失败','student13 选择数据结构','409；课表与人数均不变'],['时间冲突','选择 11 班数据结构与 3 班组成原理','409；原有课表保留'],['跨身份访问','学生访问人员维护接口','403；没有人员数据返回'],['关闭冲突','持有学期共享锁时执行关闭','409；学期保持 open'],['无教师补位','三份课表含无教师课程与相同备选','无教师班取消，备选达到 3 人开班'],['首备选满员','第一备选已 10 人','不超额，尝试第二备选'],['成绩批次回滚','先传有效成绩再传非本班学生','整个批次回滚，不保存首项'],['计费故障恢复','计费返回 503，再恢复可用','pending 保留；重试后 sent；收据不重复']]),
 h('5 执行与缺陷记录'),p('使用 tools/run-tests.sh 运行自动化检查。原始日志和 JUnit XML 存放于 03-测试证据。缺陷记录至少包括重现步骤、实际与预期结果、影响范围、修复位置和复测结果。不得用手工编写的“通过”替代命令输出。'),
 h('6 退出标准与剩余测试'),p('全部自动化用例通过；三种角色的核心页面操作可完成；启动、数据持久化和外部故障恢复有证据。500/2000 用户压力测试、80% 事务两分钟目标、长期可用性和真实 Ingres 联调另行安排；在完成前不得声明达到这些指标。')]
])

report_xml=ROOT/'03-测试证据/pytest.xml'
r=E.parse(str(report_xml));suite=r.find('.//testsuite');count=suite.get('tests');seconds=suite.get('time')
doc('05-测试报告','课程注册系统测试报告',[
 [h('1 测试结论'),p(f'2026 年 10 月 13 日的当前版本通过 {count} 项 pytest 自动化测试，失败 0 项、错误 0 项、跳过 0 项。JUnit 记录的测试套件耗时为 {seconds} 秒。前端生产构建通过，核心浏览器操作与独立 HTTP 模拟服务联调通过。当前结果支持课程演示，不构成高并发或长期可用性认证。'),
 h('2 证据与复现'),t(['证据','内容'],[['03-测试证据/pytest.txt、pytest.xml','逐项结果、运行环境、用例名称与耗时'],['03-测试证据/frontend-build.txt','Vite 生产构建输出'],['03-测试证据/docker-build.txt','Compose 构建与启动记录'],['03-测试证据/外部服务联调.txt','5 项真实 HTTP 故障与恢复断言'],['03-测试证据/浏览器检查.txt、截图','三角色界面操作及响应布局观察']]),
 h('3 关键结果'),t(['检查','观测结果'],[['并发最后名额','两个请求竞争已有 9 人的教学班，一个成功，一个返回 409；相关班级均为 10 人'],['事务回滚','冲突选课保留旧课表；成绩批次含非法学生时不保存前面的有效修改'],['权限隔离','学生不能访问师生管理或录入成绩；教师不能读取其他教师名单'],['关闭与补位','在途事务时关闭被拒绝；无教师班取消，首备选满员时使用第二备选'],['计费','关闭后生成唯一通知，故障时保留 pending，恢复后 sent，重复投递收据数量保持不变'],['界面','学生提交 4+2 后显示 4 门 13 学分；教师成功保存成绩；管理员可查看师生与模拟服务状态']]),
 h('4 测试数据隔离'),p('功能与并发用例使用 course_test，真实 HTTP 联调使用单独的临时 mock 容器。浏览器演示使用 course 中的虚构账户，保留 student01 的示例课表和演示成绩。没有使用真实学生个人信息或生产财务记录。')],
 [h('5 发现的问题及处理'),t(['问题','处理','验证'],[['Docker Hub 拉取 Node/Nginx 超时','Vue 本机预构建，Flask 托管静态文件，减少镜像依赖','最终 Compose 构建启动成功'],['中文路径作为构建上下文触发 BuildKit 头部错误','构建上下文改回 ASCII 末级目录 backend','最终构建日志无该错误'],['跨页面保留原滚动位置','导航时回到页面顶部','界面复查'],['初始演示授课存在重复时段','11 班设为待安排，冲突留作测试场景','授课冲突测试通过'],['模板中文渲染缺字','为文档渲染配置本机中文字库','逐页检查中文文本和表格']]),
 h('6 限制与未验证项目'),bullets('尚未执行题目所列 500 / 2000 同时用户压力测试，无法报告对应吞吐、P95 或错误率。','尚未进行连续多日可用性统计、故障切换、恢复时间或真实网络分区测试。','旧课程目录与计费服务为模拟实现，真实协议、身份认证和数据同步需接入后验证。','真实课程目录变化后的本地先修与费用快照同步仍需实现，演示数据固定。','个人实际贡献与工作经历由成员本人核实，不属于软件自动化测试结论。'),
 h('7 交付建议'),p('提交前在另一台装有 Docker 的机器上按 README 完整启动一次，再逐项演练学生、教师和教务操作。若源码或数据模型发生变化，重新运行相关测试并更新本报告与证据，避免报告指向旧版本。最后补齐封面中的真实组号、学号、贡献比例和本人签名。')]
])

doc('06-软件配置管理计划','软件配置管理计划',[[
 p('配置管理覆盖应用源码、依赖清单、部署文件、数据库结构、初始化数据、提交文档和测试证据。交付版本以实际源文件和 SHA256 清单为准，不伪造个人提交历史。'),
 h('1 配置项'),t(['配置项','位置','维护方式'],[['业务源码','02-源代码/backend、frontend/src','按模块修改，相关测试通过后集成'],['依赖','requirements.txt、requirements.lock.txt、package-lock.json','版本固定，依赖升级后重新构建与测试'],['部署与配置','compose.yaml、Dockerfile、.env.example','只提交示例配置；真实密钥不入库'],['数据库','models.py、schema.sql、seed.py','三者同步；已有数据不自动重置'],['文档与证据','01-提交文档、03-测试证据','与当前代码一同更新'],['原始依据','04-原始要求','保持原文件不变，作为追溯依据']]),
 h('2 版本与变更流程'),p('建议使用 main 保存可运行版本，成员通过 feature/模块名 分支开发。每次变更记录需求编号、修改文件、测试结果和风险；由 陈辰 核对后集成。提交前生成文件哈希清单。当前交付包是可导入版本控制的配置基线，版本历史由小组后续真实开发活动形成。'),
 h('3 备份与恢复'),p('数据库通过 PostgreSQL 命名卷持久化。tools/backup.sh 导出逻辑备份；恢复时先停止写入，在明确目标数据库和备份文件后使用 psql 恢复。docker compose down 不删除数据卷；只有明确需要重新初始化演示数据时才执行 down -v，该命令会删除本项目的演示数据。'),
 h('4 发布规则'),p('交付排除 node_modules、虚拟环境、临时渲染目录和私密 .env；保留已构建的 backend/static，使 Docker 演示无需安装 Node。修改 Vue 源码后先 npm ci 与 npm run build，再重建 API 镜像。冻结后修改必须更新测试结果、版本说明和 SHA256 清单。')
]])

doc('07-小组工作总结','课程注册系统小组工作总结',[[
 p('本组围绕课程注册系统完成了一个 Vue + Flask + PostgreSQL 的可运行版本，包含学生选课、教师授课及成绩、教务维护、关闭选课和模拟计费对接。交付包同时提供开发、需求、设计、测试和配置管理材料。'),
 h('1 完成情况'),p(f'当前版本已有 {count} 项自动化测试通过，包含真实 PostgreSQL 行锁并发验证和异常回滚。独立模拟服务验证了目录故障、计费持久化重试及幂等。程序通过 Docker Compose 启动，业务数据保存在命名卷中。未完成的高并发、长期可用性与真实外部系统接入已列入测试报告。'),
 h('2 分工与贡献核对'),t(['成员','建议职责','对应交付','拟定比例'],TEAMS),p('上表是按工作包划分的建议比例，五人合计 100%。实际贡献应依据成员确认的工作记录调整；本报告不能替代本人对实际工作内容的确认。最终封面必须填入确认后的比例和各成员本人签名。'),
 h('3 经验与待改进事项'),p('选课流程的难点是数据一致性：容量检查必须与写入处在同一事务中，关闭选课必须和在途修改互斥。另一重点是外部故障恢复，发送失败不能丢失计费信息，重试必须避免重复记账。后续可完善数据库迁移、真实目录快照同步、登录限流及规模化性能验证。'),
 h('4 提交前行动'),p('由 陈辰 汇总五位成员核实后的个人总结，补齐组号与学号，确定实际贡献比例并组织签名。全组共同进行一轮现场演示，确认提交文档、源码及运行结果对应同一版本。')
]])

PERSONAL=[
 ('组长与系统集成','开发计划、总体架构、Compose 编排与全组总结','app.py 的系统边界、compose.yaml、开发计划及小组总结','核对需求范围和交付清单，检查接口及数据结构的一致性，组织测试、演示和版本整理。','关闭选课需要与在途注册事务互斥；工作分配应以可验收功能为单位，避免文档与源码分别演进。'),
 ('需求分析与学生端','需求分析、课程目录、选课清单、课表及成绩页面','frontend/src/App.vue 的学生工作台、需求分析文档','区分课程和教学班，明确草稿、已提交、备选、取消等状态，逐条对应学生用例和异常提示。','主选和备选的业务含义不同，界面必须清楚说明备选不占名额；提交失败时必须保留学生原来的课表。'),
 ('后端业务与数据库','关系模型、选课事务、并发锁、容量和先修校验','backend/models.py、app.py、schema.sql','检查数据库约束和接口校验，验证人数边界、冲突回滚、关闭事务和费用计算的一致性。','先查询容量再独立写入会产生并发超额。需要锁定共同的教学班记录，在事务内校验并提交。'),
 ('教师与教务模块','授课安排、名单、成绩录入、人员维护和授课资格','frontend/src/App.vue 的教师/教务页面、app.py 对应接口','核对任课归属、教师资格、成绩等级及人员字段；验证不属于本人教学班的请求无法获取或修改数据。','隐藏前端按钮不能代替后端权限控制。成绩批量保存需要原子性，避免一个非法学生导致部分成绩被悄悄修改。'),
 ('测试部署与交付','测试计划、自动化用例、外部联调、证据及使用说明','backend/tests、check_integrations.py、worker.py 的验证、README 与测试报告','使用独立测试数据库复现关键场景，保留原始命令输出，核对运行步骤、数据卷及失败恢复行为。','测试应同时验证返回值和数据库状态。模拟接口通过不等于真实接口通过，小规模并发验证也不能替代 2000 用户压力测试。'),
]
for i,(role,work,files,steps,lesson) in enumerate(PERSONAL):
    doc(f'08-{i}-{MEMBER_NAMES[i]}个人工作总结',f'{MEMBER_NAMES[i]}个人工作总结',[[
        p(f'姓名：{MEMBER_NAMES[i]}。'+('学号：25230426。' if i==0 else '学号：14230107。' if i==2 else '')+f'建议职责：{role}。本报告按项目工作包整理，提交前由本人核对实际参与内容、补充个人工作记录并确认贡献比例。'),
        h('1 对应工作范围'),p(work+'。'),h('2 对应成果'),p(files+'。'),
        h('3 工作核对说明'),p(steps),p('请本人补充实际完成的功能、遇到的问题、参与日期、修改记录和测试证据。当前文档描述模块成果，不将代码自动生成过程记作任何成员已经完成的个人经历。'),
        h('4 技术认识与总结'),p(lesson),h('5 后续工作'),p('结合本人负责模块运行相关测试，复查需求、设计和操作说明，参加全组演示。若发现文档与程序行为不一致，先修正并复测，再提交最终个人总结。'),
        h('6 贡献与签字'),p('建议工作量比例：20%。实际贡献比例：____________。'),p('本人签名：____________________    日期：____________________')
    ]])

def set_font(run,size=11,bold=False):
    run.font.name='Songti SC';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')

def make_table(d,heads,data):
    tb=d.add_table(rows=1,cols=len(heads));tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    tb.autofit=False
    widths={2:[4.1,12.5],3:[3.1,6.7,6.8],4:[2.6,3.0,8.3,2.7]}[len(heads)]
    if heads==['编号','功能','验收要点']:widths=[1.4,3.3,11.9]
    if heads==['方法与路径','角色','含义']:widths=[7.8,2.2,6.6]
    for i,w in enumerate(widths):tb.columns[i].width=Cm(w)
    for i,x in enumerate(heads):tb.rows[0].cells[i].text=x
    for row in data:
        cells=tb.add_row().cells
        for cell,value in zip(cells,row):cell.text=str(value)
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders);m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');m.append(el)
            pr.append(m)
            if ri==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EEF0');pr.append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=Pt(13)
                for run in para.runs:set_font(run,9.5,ri==0)
    spacer=d.add_paragraph()
    spacer.paragraph_format.space_after=Pt(0)
    spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(6)
    spacer.add_run(' ').font.size=Pt(1)

for name,title,pages in DOCS:
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2.0);s.bottom_margin=Cm(2.0);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
    for sty in ['Normal','Title','Heading 1','Heading 2','List Bullet']:
        st=d.styles[sty];st.font.name='Songti SC';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
    normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=Pt(16);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.space_before=Pt(0)
    d.styles['Title'].font.size=Pt(21);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.space_after=Pt(9)
    d.styles['Heading 1'].font.size=Pt(13);d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(12);d.styles['Heading 1'].paragraph_format.space_after=Pt(6)
    d.styles['Title'].paragraph_format.line_spacing=Pt(27)
    d.styles['Heading 1'].paragraph_format.line_spacing=Pt(19)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for grid in list(s._sectPr.iter(qn('w:docGrid'))):grid.getparent().remove(grid)
    d.add_paragraph(title,'Title');d.add_paragraph('2026 软件工程课程设计    陈辰 小组    版本 1.0    2026-10-13').runs[0].font.size=Pt(9)
    for page_index,blocks in enumerate(pages):
        # Use pageBreakBefore on the next heading to avoid empty spill pages.
        for block in blocks:
            if block[0]=='h':
                para=d.add_paragraph(block[1],'Heading 1')
                if page_index and block is blocks[0]:para.paragraph_format.page_break_before=True
            elif block[0]=='p':d.add_paragraph(block[1])
            elif block[0]=='b':
                for item in block[1]:d.add_paragraph(item,'List Bullet')
            elif block[0]=='t':make_table(d,block[1],block[2])
    for para in d.element.body.iter(qn('w:p')):
        pp=para.find(qn('w:pPr'))
        if pp is None:pp=E.SubElement(para,qn('w:pPr'))
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run();set_font(r,9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    d.core_properties.title=title;d.core_properties.author='陈辰 小组';d.save(OUT/(name+'.docx'))
print('Created',len(DOCS)+1,'DOCX files')
